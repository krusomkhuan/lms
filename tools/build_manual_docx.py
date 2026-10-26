"""สร้างคู่มือการใช้งานเป็นไฟล์ Word (.docx) จาก JSON ที่ export_manual.php ส่งออก

ใช้: python tools/build_manual_docx.py manual.json docs/manual-vocational-teacher.docx

- ฟอนต์ TH Sarabun New 16pt, A4 ระยะขอบตามแบบหนังสือราชการ (บน 2.5 ล่าง 2 ซ้าย 3 ขวา 2 ซม.)
- แทรก Zero-Width Space ระหว่างคำภาษาไทย (pythainlp) ให้ Word ตัดบรรทัดไทยได้เต็มบรรทัด
- ภาษาไทยเป็น complex script ใน Word จึงต้องตั้ง bCs / szCs คู่กับ b / sz เสมอ
"""
import json
import re
import sys
from html.parser import HTMLParser

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pythainlp.tokenize import word_tokenize

FONT = "TH Sarabun New"
BODY_PT = 16
NAVY = RGBColor(0x0D, 0x3B, 0x66)
TEAL = RGBColor(0x1B, 0x99, 0x8B)
ZWSP = "​"
THAI_RUN = re.compile(r"[฀-๿]+")

# ไอคอน Bootstrap ในเนื้อหาเว็บ → ข้อความในเอกสาร (ไอคอนประดับที่เหลือตัดทิ้ง)
ICON_TEXT = {
    "bi-list": "☰", "bi-check-lg": "✓", "bi-slash-circle": "⊘", "bi-trash": "[ลบ]",
    "bi-download": "[ดาวน์โหลด]", "bi-arrow-counterclockwise": "[กู้คืน]", "bi-stars": "[AI]",
    "bi-collection-play": "[บทเรียน]", "bi-people": "[ผู้เรียน]",
}
ALERT_LABEL = {"info": ("เคล็ดลับ", "E7F1FB", "2F6FB0"), "success": ("แนะนำ", "E6F4EA", "2E7D4F"),
               "warning": ("ข้อควรระวัง", "FFF4E0", "C77700"), "secondary": ("หมายเหตุ", "F1F3F5", "6C757D")}


def insert_zwsp(text: str) -> str:
    """แทรก ZWS ระหว่างคำไทย เฉพาะช่วงที่เป็นอักษรไทย (อังกฤษ/ตัวเลข/URL ไม่ถูกแตะ)"""
    return THAI_RUN.sub(lambda m: ZWSP.join(w for w in word_tokenize(m.group(0), engine="newmm") if w), text)


# ---------- HTML → ต้นไม้อย่างง่าย ----------
class Node:
    def __init__(self, tag, attrs=None, parent=None):
        self.tag, self.attrs, self.parent, self.children = tag, dict(attrs or {}), parent, []

    def cls(self):
        return self.attrs.get("class", "").split()


class TreeBuilder(HTMLParser):
    VOID = {"br", "hr", "img"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("root")
        self.cur = self.root

    def handle_starttag(self, tag, attrs):
        n = Node(tag, attrs, self.cur)
        self.cur.children.append(n)
        if tag not in self.VOID:
            self.cur = n

    def handle_endtag(self, tag):
        n = self.cur
        while n is not self.root and n.tag != tag:
            n = n.parent
        if n is not self.root:
            self.cur = n.parent

    def handle_data(self, data):
        self.cur.children.append(data)


def parse(html):
    b = TreeBuilder()
    b.feed(html)
    return b.root


# ---------- รูปแบบตัวอักษร (รองรับภาษาไทย) ----------
def set_run_font(run, size=None, bold=None, color=None, underline=None):
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    if bold is not None:
        run.bold = bold
        bcs = rpr.find(qn("w:bCs"))
        if bcs is None:
            bcs = OxmlElement("w:bCs")
            rpr.append(bcs)
        bcs.set(qn("w:val"), "1" if bold else "0")
    if size is not None:
        run.font.size = Pt(size)
        szcs = rpr.find(qn("w:szCs"))
        if szcs is None:
            szcs = OxmlElement("w:szCs")
            rpr.append(szcs)
        szcs.set(qn("w:val"), str(int(size * 2)))
    if color is not None:
        run.font.color.rgb = color
    if underline is not None:
        run.underline = underline


def style_fonts(style, size, bold=False, color=None):
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(a), FONT)
    for theme in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if fonts.get(qn(theme)) is not None:
            del fonts.attrib[qn(theme)]
    style.font.size = Pt(size)
    style.font.bold = bold
    for tag, val in (("w:szCs", str(size * 2)), ("w:bCs", "1" if bold else "0")):
        el = rpr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr.append(el)
        el.set(qn("w:val"), val)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:bidi"), "th-TH")
    if color is not None:
        style.font.color.rgb = color


def add_text(par, text, bold=False, underline=False, color=None, size=None):
    if not text:
        return
    run = par.add_run(insert_zwsp(text))
    set_run_font(run, size=size, bold=bold if bold else None, color=color, underline=underline or None)


def add_inline(par, node, fmt=None, size=None):
    """เพิ่มเนื้อหาแบบ inline (ข้อความ, ตัวหนา, ลิงก์, ไอคอน) ลงในย่อหน้า"""
    fmt = fmt or {}
    for ch in node.children:
        if isinstance(ch, str):
            add_text(par, re.sub(r"\s+", " ", ch), size=size, **fmt)
            continue
        if ch.tag in ("ul", "ol"):
            continue                                     # รายการซ้อน จัดการแยก
        if ch.tag == "i":
            icon = next((c for c in ch.cls() if c.startswith("bi-")), "")
            if icon in ICON_TEXT:
                add_text(par, ICON_TEXT[icon], bold=True, color=NAVY, size=size)
            continue
        if ch.tag == "br":
            par.add_run().add_break()
            continue
        f = dict(fmt)
        if ch.tag in ("b", "strong") or "badge" in ch.cls():
            f["bold"] = True
        if ch.tag == "u":
            f["underline"] = True
        if ch.tag == "a":
            f["color"] = NAVY
        add_inline(par, ch, f, size)


def trim(par):
    """ตัดช่องว่างหัว/ท้ายย่อหน้า"""
    runs = par.runs
    if runs:
        runs[0].text = runs[0].text.lstrip()
        runs[-1].text = runs[-1].text.rstrip()


def list_paragraph(doc, prefix, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.9 + 0.8 * level)
    pf.first_line_indent = Cm(-0.7)
    pf.space_after = Pt(2)
    pf.tab_stops.add_tab_stop(Cm(0.9 + 0.8 * level))
    r = p.add_run(prefix + "\t")
    set_run_font(r, bold=True, color=TEAL if prefix[0].isdigit() else NAVY)
    return p


def render_list(doc, node, level=0):
    ordered = node.tag == "ol"
    n = 0
    for li in node.children:
        if isinstance(li, str) or li.tag != "li":
            continue
        n += 1
        p = list_paragraph(doc, f"{n}." if ordered else "•", level)
        add_inline(p, li)
        trim(p)
        for sub in li.children:
            if not isinstance(sub, str) and sub.tag in ("ul", "ol"):
                render_list(doc, sub, level + 1)


def shade_cell(cell, fill, border_color):
    tcpr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "nil")
        borders.append(b)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:color"), border_color)
    borders.append(left)
    tcpr.append(borders)


def render_alert(doc, node):
    kind = next((c[6:] for c in node.cls() if c.startswith("alert-")), "info")
    label, fill, color = ALERT_LABEL.get(kind, ALERT_LABEL["info"])
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    shade_cell(cell, fill, color)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    add_text(p, label + ": ", bold=True, color=RGBColor.from_string(color), size=BODY_PT - 1)
    add_inline(p, node, size=BODY_PT - 1)
    trim(p)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def render_html(doc, html):
    for ch in parse(html).children:
        if isinstance(ch, str):
            if ch.strip():
                p = doc.add_paragraph()
                add_text(p, ch.strip())
            continue
        if ch.tag in ("ol", "ul"):
            render_list(doc, ch)
        elif ch.tag == "div" and "alert" in ch.cls():
            render_alert(doc, ch)
        elif ch.tag == "h6":
            p = doc.add_paragraph(style="Heading 3")
            add_inline(p, ch)
            trim(p)
        else:
            p = doc.add_paragraph()
            add_inline(p, ch)
            trim(p)


# ---------- ฟิลด์ (สารบัญ / เลขหน้า) ----------
def add_field(par, instr, placeholder=""):
    def fld(kind):
        r = par.add_run()
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), kind)
        r._element.append(el)
        return r
    fld("begin")
    r = par.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r._element.append(it)
    fld("separate")
    if placeholder:
        add_text(par, placeholder)
    fld("end")


def build(data, out):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Cm(2.5), Cm(2), Cm(3), Cm(2)

    styles = doc.styles
    style_fonts(styles["Normal"], BODY_PT)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.line_spacing = 1.0
    style_fonts(styles["Heading 1"], 24, True, NAVY)
    style_fonts(styles["Heading 2"], 19, True, NAVY)
    style_fonts(styles["Heading 3"], 17, True, TEAL)
    styles["Heading 1"].paragraph_format.page_break_before = True
    styles["Heading 1"].paragraph_format.space_after = Pt(10)
    styles["Heading 2"].paragraph_format.space_before = Pt(14)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    for name in ("TOC 1", "TOC 2"):
        try:
            style_fonts(styles[name], BODY_PT, name == "TOC 1")
        except KeyError:
            pass

    # ----- หน้าปก -----
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "คู่มือการใช้งานระบบ", bold=True, color=NAVY, size=40)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, data["site_name"], bold=True, color=TEAL, size=30)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, data["slogan"], size=22)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "ระบบจัดการเรียนการสอนออนไลน์ (LMS) สำหรับครูอาชีวศึกษา", size=18)
    for _ in range(8):
        doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, "สำหรับผู้เรียน · ครูผู้สอน · ผู้ดูแลระบบ", size=18)
    if data.get("org"):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, data["org"], bold=True, size=18)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(p, f"เวอร์ชันระบบ {data['version']}  ·  {data['date_th']}", size=16, color=RGBColor(0x55, 0x55, 0x55))

    # ----- สารบัญ (Word อัปเดตเลขหน้าตอนแปลงเป็น PDF) -----
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    add_text(p, "สารบัญ", bold=True, color=NAVY, size=24)
    p = doc.add_paragraph()
    add_field(p, 'TOC \\o "1-2" \\h \\z \\u', "(คลิกขวา → Update Field เพื่อแสดงสารบัญ)")

    # ----- หัว/ท้ายกระดาษ -----
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_text(header, f"คู่มือการใช้งาน {data['site_name']} v{data['version']}", size=12, color=RGBColor(0x88, 0x88, 0x88))
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "หน้า ", size=13)
    add_field(footer, "PAGE", "1")
    sec.different_first_page_header_footer = True     # หน้าปกไม่มีหัว/ท้ายกระดาษ

    # ----- เนื้อหา -----
    titles = {"student": "สำหรับผู้เรียน", "teacher": "สำหรับครูผู้สอน", "admin": "สำหรับผู้ดูแลระบบ", "faq": "คำถามที่พบบ่อย"}
    for pi, part in enumerate(data["parts"], 1):
        h = doc.add_paragraph(style="Heading 1")
        add_text(h, f"ส่วนที่ {pi}  {titles.get(part['key'], part['label'])}")
        for si, s in enumerate(part["sections"], 1):
            h = doc.add_paragraph(style="Heading 2")
            add_text(h, f"{pi}.{si}  {s['title']}")
            render_html(doc, s["html"])

    doc.core_properties.title = f"คู่มือการใช้งาน {data['site_name']}"
    doc.core_properties.subject = data["slogan"]
    doc.core_properties.author = data["site_name"]
    doc.core_properties.language = "th-TH"
    doc.save(out)


if __name__ == "__main__":
    with open(sys.argv[1], encoding="utf-8") as f:
        build(json.load(f), sys.argv[2])
    print("saved", sys.argv[2])
